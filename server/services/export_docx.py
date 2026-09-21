"""一键导出 Word（python-docx）：封面 / 目录页 / 卷章层级 / 页眉页脚 / 段首缩进。

- 导出前经排版引擎（复用 typesetting.normalize），并把 rules_version 写入文档属性，
  保证「导出与预览一致」。
- export_templates 保存多套样式。
"""
import json
from pathlib import Path

from server.config import DATA_DIR
from server.db import get_conn
from server.models import books as bm, chapters as chm
from server.services import typesetting as ts

EXPORT_DIR = DATA_DIR / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

DEFAULT_STYLE = {
    "font": "宋体",
    "size": 12,
    "line_spacing": 1.5,
    "indent_chars": 2,
    "cover": True,
    "toc": True,
    "header": True,
    "footer": True,
}


# ---------- 模板 ----------
def list_templates():
    conn = get_conn()
    try:
        return [dict(r) for r in conn.execute("SELECT * FROM export_templates ORDER BY built_in DESC, id").fetchall()]
    finally:
        conn.close()


def create_template(name, style):
    conn = get_conn()
    try:
        cur = conn.execute(
            "INSERT INTO export_templates(name, style_json) VALUES (?,?)",
            (name, json.dumps(style, ensure_ascii=False)),
        )
        conn.commit()
        return dict(conn.execute("SELECT * FROM export_templates WHERE id=?", (cur.lastrowid,)).fetchone())
    finally:
        conn.close()


def _style(style=None):
    s = dict(DEFAULT_STYLE)
    s.update(style or {})
    return s


def _add_para(doc, text, font, size):
    """新增段落并写入一个已设字体的 run。"""
    p = doc.add_paragraph()
    _set_font(p.add_run(text or ""), font, size)
    return p


def _set_font(run, font, size):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = font
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def _add_page_number(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = "PAGE"
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)


def export_book(book_id, style=None, typeset=True, filename=None):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    st = _style(style)
    book = bm.get_book(book_id)
    if not book:
        raise ValueError("作品不存在")
    rules = ts.load_rules()

    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = st["font"]
    normal.font.size = Pt(st["size"])
    try:
        from docx.oxml.ns import qn
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), st["font"])
    except Exception:
        pass

    # 封面
    if st.get("cover"):
        for _ in range(6):
            doc.add_paragraph("")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.add_run(book["title"]), st["font"], st["size"] + 16)
        if book.get("author"):
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_font(p2.add_run(book["author"]), st["font"], st["size"] + 4)
        doc.add_page_break()

    # 目录页
    volumes = bm.list_volumes(book_id)
    rows = chm.list_chapters(book_id)
    if st.get("toc"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.add_run("目录"), st["font"], st["size"] + 8)
        for v in volumes:
            _add_para(doc, v["title"], st["font"], st["size"])
            for c in rows:
                if (c.get("volume_id") or None) == v["id"]:
                    _add_para(doc, "    " + (c.get("title") or ""), st["font"], st["size"])
        for c in rows:
            if not c.get("volume_id"):
                _add_para(doc, c.get("title") or "", st["font"], st["size"])
        doc.add_page_break()

    # 页眉页脚
    sec = doc.sections[0]
    if st.get("header"):
        hp = sec.header.paragraphs[0]
        hp.text = book["title"]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if st.get("footer"):
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number(fp)

    # 正文：卷 → 章
    written = 0

    def add_chapter(c):
        nonlocal written
        full = chm.get_chapter(c["id"])
        content = (full.get("content") or "")
        if typeset and content:
            content = ts.normalize(content, {})["content"]
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.add_run(c.get("title") or "未命名章节"), st["font"], st["size"] + 6)
        for para in content.split("\n"):
            if not para.strip():
                continue
            bp = doc.add_paragraph()
            bp.paragraph_format.line_spacing = st["line_spacing"]
            bp.paragraph_format.first_line_indent = Pt(st["size"] * int(st["indent_chars"]))
            _set_font(bp.add_run(para.strip()), st["font"], st["size"])
        doc.add_page_break()
        written += 1

    for v in volumes:
        vp = doc.add_paragraph()
        vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(vp.add_run(v["title"]), st["font"], st["size"] + 10)
        doc.add_page_break()
        for c in rows:
            if (c.get("volume_id") or None) == v["id"]:
                add_chapter(c)
    for c in rows:
        if not c.get("volume_id"):
            add_chapter(c)

    # 记录排版规则版本，保证与预览一致
    doc.core_properties.comments = f"rules_version={rules.get('rules_version')}"
    doc.core_properties.title = book["title"]

    name = filename or f"{book['title']}.docx"
    out = EXPORT_DIR / name
    doc.save(str(out))
    return {"path": str(out), "bytes": out.stat().st_size, "chapters": written,
            "rules_version": rules.get("rules_version")}
